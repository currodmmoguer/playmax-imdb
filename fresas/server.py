"""Fresas — gestión de una explotación de fresa al estilo de Huelva.

Backend Flask + SQLite para llevar la campaña: fincas y sectores
(macrotúneles), trabajadores, recolección diaria, entregas a la cooperativa
y el cuaderno de campo de tratamientos fitosanitarios (con control del plazo
de seguridad antes de recolectar).

El acceso requiere usuario y contraseña. El primer administrador se crea
desde la consola:  python server.py crear-admin <usuario>
"""

import collections
import datetime
import getpass
import io
import os
import secrets
import sqlite3
import sys
import threading
import time

from flask import (Flask, g, jsonify, request, send_file, send_from_directory,
                   session)
from openpyxl import Workbook
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
DB_PATH = os.environ.get("FRESAS_DB", os.path.join(DATA_DIR, "fresas.db"))
SECRET_KEY_FILE = os.path.join(DATA_DIR, "secret_key")

# En producción la app va detrás de un proxy HTTPS (Tailscale Funnel,
# Cloudflare Tunnel, Caddy…): FRESAS_HTTPS=1 marca la cookie de sesión como
# Secure y FRESAS_PROXY=1 confía en las cabeceras X-Forwarded-* de ese proxy.
HTTPS = os.environ.get("FRESAS_HTTPS") == "1"
BEHIND_PROXY = os.environ.get("FRESAS_PROXY") == "1"


def load_secret_key():
    if os.environ.get("FRESAS_SECRET_KEY"):
        return os.environ["FRESAS_SECRET_KEY"]
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(SECRET_KEY_FILE):
        fd = os.open(SECRET_KEY_FILE, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "w") as f:
            f.write(secrets.token_hex(32))
    with open(SECRET_KEY_FILE) as f:
        return f.read().strip()


app = Flask(__name__, static_folder="static", static_url_path="")
app.config.update(
    SEND_FILE_MAX_AGE_DEFAULT=0,
    SECRET_KEY=load_secret_key(),
    SESSION_COOKIE_NAME="fresas_sesion",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=HTTPS,
    # En el campo se entra desde el móvil: la sesión dura 30 días.
    PERMANENT_SESSION_LIFETIME=datetime.timedelta(days=30),
)
if BEHIND_PROXY:
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

ROLES = ["admin", "usuario"]
MIN_PASSWORD = 8

CATEGORIAS = ["Extra", "Primera", "Segunda", "Industria"]

SCHEMA = """
CREATE TABLE IF NOT EXISTS fincas (
    id INTEGER PRIMARY KEY,
    nombre TEXT NOT NULL,
    municipio TEXT,
    notas TEXT
);
CREATE TABLE IF NOT EXISTS sectores (
    id INTEGER PRIMARY KEY,
    finca_id INTEGER NOT NULL REFERENCES fincas(id),
    nombre TEXT NOT NULL,
    variedad TEXT,
    superficie_ha REAL,
    plantas INTEGER,
    fecha_plantacion TEXT
);
CREATE TABLE IF NOT EXISTS trabajadores (
    id INTEGER PRIMARY KEY,
    nombre TEXT NOT NULL,
    dni TEXT,
    telefono TEXT,
    cuadrilla TEXT,
    fecha_alta TEXT,
    activo INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS recolecciones (
    id INTEGER PRIMARY KEY,
    fecha TEXT NOT NULL,
    sector_id INTEGER NOT NULL REFERENCES sectores(id),
    trabajador_id INTEGER REFERENCES trabajadores(id),
    categoria TEXT NOT NULL,
    cajas INTEGER,
    kg REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS entregas (
    id INTEGER PRIMARY KEY,
    fecha TEXT NOT NULL,
    destino TEXT NOT NULL,
    albaran TEXT,
    categoria TEXT NOT NULL,
    kg REAL NOT NULL,
    precio_kg REAL
);
CREATE TABLE IF NOT EXISTS tratamientos (
    id INTEGER PRIMARY KEY,
    fecha TEXT NOT NULL,
    sector_id INTEGER NOT NULL REFERENCES sectores(id),
    producto TEXT NOT NULL,
    num_registro TEXT,
    materia_activa TEXT,
    plaga TEXT,
    dosis TEXT,
    plazo_seguridad_dias INTEGER NOT NULL DEFAULT 0,
    aplicador TEXT
);
CREATE INDEX IF NOT EXISTS idx_recol_fecha ON recolecciones(fecha);
CREATE INDEX IF NOT EXISTS idx_entregas_fecha ON entregas(fecha);
CREATE INDEX IF NOT EXISTS idx_trat_sector ON tratamientos(sector_id, fecha);
CREATE TABLE IF NOT EXISTS usuarios (
    id INTEGER PRIMARY KEY,
    usuario TEXT NOT NULL UNIQUE COLLATE NOCASE,
    nombre TEXT,
    password_hash TEXT NOT NULL,
    rol TEXT NOT NULL DEFAULT 'usuario',
    activo INTEGER NOT NULL DEFAULT 1,
    creado TEXT NOT NULL DEFAULT (datetime('now'))
);
"""

# Definición de las tablas editables desde la API genérica: campo -> tipo.
TABLES = {
    "fincas": {
        "fields": {"nombre": "str", "municipio": "str", "notas": "str"},
        "required": ["nombre"],
        "order": "nombre",
    },
    "sectores": {
        "fields": {"finca_id": "int", "nombre": "str", "variedad": "str",
                   "superficie_ha": "float", "plantas": "int",
                   "fecha_plantacion": "date"},
        "required": ["finca_id", "nombre"],
        "order": "finca_id, nombre",
    },
    "trabajadores": {
        "fields": {"nombre": "str", "dni": "str", "telefono": "str",
                   "cuadrilla": "str", "fecha_alta": "date", "activo": "bool"},
        "required": ["nombre"],
        "order": "activo DESC, nombre",
    },
    "recolecciones": {
        "fields": {"fecha": "date", "sector_id": "int", "trabajador_id": "int",
                   "categoria": "categoria", "cajas": "int", "kg": "float"},
        "required": ["fecha", "sector_id", "categoria", "kg"],
        "order": "fecha DESC, id DESC",
        "campana": True,
    },
    "entregas": {
        "fields": {"fecha": "date", "destino": "str", "albaran": "str",
                   "categoria": "categoria", "kg": "float", "precio_kg": "float"},
        "required": ["fecha", "destino", "categoria", "kg"],
        "order": "fecha DESC, id DESC",
        "campana": True,
    },
    "tratamientos": {
        "fields": {"fecha": "date", "sector_id": "int", "producto": "str",
                   "num_registro": "str", "materia_activa": "str",
                   "plaga": "str", "dosis": "str",
                   "plazo_seguridad_dias": "int", "aplicador": "str"},
        "required": ["fecha", "sector_id", "producto"],
        "order": "fecha DESC, id DESC",
        "campana": True,
    },
}


# ---------------------------------------------------------------------------
# Base de datos
# ---------------------------------------------------------------------------

def campana_de(fecha):
    """Campaña de una fecha ISO. En Huelva se planta en octubre y se recoge
    hasta junio, así que la campaña va de septiembre a agosto: 2026/27."""
    if not fecha:
        return None
    y, m = int(fecha[:4]), int(fecha[5:7])
    if m < 9:
        y -= 1
    return f"{y}/{(y + 1) % 100:02d}"


def get_db():
    if "db" not in g:
        os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
        db = sqlite3.connect(DB_PATH)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys = ON")
        db.create_function("campana", 1, campana_de, deterministic=True)
        g.db = db
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    with app.app_context():
        get_db().executescript(SCHEMA)


def rows(cur):
    return [dict(r) for r in cur.fetchall()]


class ValidationError(Exception):
    pass


def parse_payload(table, data):
    spec = TABLES[table]
    out = {}
    for name, kind in spec["fields"].items():
        if name not in data:
            continue
        v = data[name]
        if v in ("", None):
            out[name] = None
            continue
        try:
            if kind == "int":
                v = int(v)
            elif kind == "float":
                v = float(str(v).replace(",", "."))
            elif kind == "bool":
                v = 1 if v in (True, 1, "1", "true", "on") else 0
            elif kind == "date":
                v = datetime.date.fromisoformat(str(v)).isoformat()
            elif kind == "categoria":
                if v not in CATEGORIAS:
                    raise ValueError
            else:
                v = str(v).strip()
        except ValueError:
            raise ValidationError(f"Valor no válido en «{name}»")
        if kind in ("int", "float") and v < 0:
            raise ValidationError(f"«{name}» no puede ser negativo")
        out[name] = v
    return out


def check_required(table, values):
    missing = [f for f in TABLES[table]["required"] if values.get(f) is None]
    if missing:
        raise ValidationError("Faltan campos: " + ", ".join(missing))


def tratamientos_activos(db, sector_id, fecha):
    """Tratamientos cuyo plazo de seguridad todavía no ha vencido en `fecha`."""
    return rows(db.execute(
        """SELECT t.*, date(t.fecha, '+' || t.plazo_seguridad_dias || ' days')
                  AS recolectable_desde
           FROM tratamientos t
           WHERE t.sector_id = ? AND t.fecha <= ?
             AND date(t.fecha, '+' || t.plazo_seguridad_dias || ' days') > ?
           ORDER BY recolectable_desde DESC""",
        (sector_id, fecha, fecha)))


# ---------------------------------------------------------------------------
# Usuarios y sesión
# ---------------------------------------------------------------------------

# Rutas accesibles sin sesión: la propia pantalla (no contiene datos) y el login.
PUBLIC_PATHS = {"/", "/index.html", "/app.js", "/style.css", "/favicon.svg",
                "/api/login"}

# Freno a ataques de fuerza bruta: como mucho 5 intentos fallidos por IP y
# usuario cada 15 minutos.
_fallos = collections.defaultdict(collections.deque)
_fallos_lock = threading.Lock()
MAX_FALLOS, VENTANA_FALLOS = 5, 15 * 60
# Hash de relleno para que tarde lo mismo si el usuario no existe.
_DUMMY_HASH = generate_password_hash(secrets.token_hex(16))


def _fallos_recientes(clave):
    q = _fallos[clave]
    now = time.time()
    while q and now - q[0] > VENTANA_FALLOS:
        q.popleft()
    return q


def check_password_rules(password):
    if not isinstance(password, str) or len(password) < MIN_PASSWORD:
        raise ValidationError(
            f"La contraseña debe tener al menos {MIN_PASSWORD} caracteres")


def password_version(pw_hash):
    # Se guarda en la sesión: al cambiar la contraseña se cierran las demás
    # sesiones abiertas de ese usuario.
    return pw_hash[-12:]


def public_user(u):
    return {"id": u["id"], "usuario": u["usuario"], "nombre": u["nombre"],
            "rol": u["rol"], "activo": bool(u["activo"]), "creado": u["creado"]}


def current_user():
    if "user" not in g:
        g.user = None
        uid = session.get("uid")
        if uid is not None:
            u = get_db().execute("SELECT * FROM usuarios WHERE id = ?",
                                 (uid,)).fetchone()
            if (u and u["activo"]
                    and session.get("ver") == password_version(u["password_hash"])):
                g.user = u
    return g.user


@app.before_request
def require_login():
    if request.path in PUBLIC_PATHS:
        return None
    # Protección CSRF: las peticiones que modifican datos deben llevar una
    # cabecera propia, que un formulario de otra web no puede añadir.
    if (request.method not in ("GET", "HEAD", "OPTIONS")
            and request.headers.get("X-Requested-With") != "fresas"):
        return jsonify(error="Petición no permitida"), 403
    if current_user() is None:
        return jsonify(error="Sesión no iniciada", login=True), 401
    if request.path.startswith("/api/usuarios") and current_user()["rol"] != "admin":
        return jsonify(error="Solo para administradores"), 403
    return None


@app.after_request
def security_headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data:; frame-ancestors 'none'; form-action 'self'")
    if HTTPS:
        resp.headers["Strict-Transport-Security"] = "max-age=31536000"
    if request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/api/login", methods=["POST"])
def login():
    if request.headers.get("X-Requested-With") != "fresas":
        return jsonify(error="Petición no permitida"), 403
    data = request.get_json(force=True, silent=True) or {}
    usuario = str(data.get("usuario", "")).strip()
    password = str(data.get("password", ""))
    clave = (request.remote_addr, usuario.lower())
    with _fallos_lock:
        if len(_fallos_recientes(clave)) >= MAX_FALLOS:
            return jsonify(error="Demasiados intentos. Espera unos minutos."), 429
    u = get_db().execute("SELECT * FROM usuarios WHERE usuario = ?",
                         (usuario,)).fetchone()
    ok = check_password_hash(u["password_hash"] if u else _DUMMY_HASH, password)
    if not (u and ok and u["activo"]):
        with _fallos_lock:
            _fallos_recientes(clave).append(time.time())
            if len(_fallos) > 10000:
                for k in [k for k in _fallos if not _fallos_recientes(k)]:
                    del _fallos[k]
        return jsonify(error="Usuario o contraseña incorrectos"), 401
    with _fallos_lock:
        _fallos.pop(clave, None)
    session.clear()
    session.permanent = True
    session["uid"] = u["id"]
    session["ver"] = password_version(u["password_hash"])
    return jsonify(public_user(u))


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify(ok=True)


@app.route("/api/me")
def me():
    return jsonify(public_user(current_user()))


@app.route("/api/me/password", methods=["POST"])
def change_own_password():
    data = request.get_json(force=True, silent=True) or {}
    u = current_user()
    if not check_password_hash(u["password_hash"], str(data.get("actual", ""))):
        return jsonify(error="La contraseña actual no es correcta"), 400
    check_password_rules(data.get("nueva"))
    pw_hash = generate_password_hash(data["nueva"])
    db = get_db()
    db.execute("UPDATE usuarios SET password_hash = ? WHERE id = ?", (pw_hash, u["id"]))
    db.commit()
    session["ver"] = password_version(pw_hash)
    return jsonify(ok=True)


@app.route("/api/usuarios")
def list_usuarios():
    return jsonify([public_user(u) for u in get_db().execute(
        "SELECT * FROM usuarios ORDER BY activo DESC, usuario")])


def parse_usuario(data, nuevo):
    out = {}
    if nuevo or "usuario" in data:
        usuario = str(data.get("usuario", "")).strip()
        if not usuario or len(usuario) > 40 or any(c.isspace() for c in usuario):
            raise ValidationError("Nombre de usuario no válido (sin espacios)")
        out["usuario"] = usuario
    if "nombre" in data:
        out["nombre"] = str(data["nombre"] or "").strip() or None
    if "rol" in data:
        if data["rol"] not in ROLES:
            raise ValidationError("Rol no válido")
        out["rol"] = data["rol"]
    if "activo" in data:
        out["activo"] = 1 if data["activo"] else 0
    if nuevo or data.get("password"):
        check_password_rules(data.get("password"))
        out["password_hash"] = generate_password_hash(data["password"])
    return out


@app.route("/api/usuarios", methods=["POST"])
def create_usuario():
    values = parse_usuario(request.get_json(force=True, silent=True) or {}, True)
    db = get_db()
    try:
        cur = db.execute(
            f"INSERT INTO usuarios ({', '.join(values)}) "
            f"VALUES ({', '.join('?' for _ in values)})", list(values.values()))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify(error="Ese nombre de usuario ya existe"), 409
    return jsonify(public_user(db.execute(
        "SELECT * FROM usuarios WHERE id = ?", (cur.lastrowid,)).fetchone())), 201


@app.route("/api/usuarios/<int:uid>", methods=["PUT"])
def update_usuario(uid):
    db = get_db()
    if db.execute("SELECT 1 FROM usuarios WHERE id = ?", (uid,)).fetchone() is None:
        return jsonify(error="No existe"), 404
    values = parse_usuario(request.get_json(force=True, silent=True) or {}, False)
    if not values:
        return jsonify(error="Nada que cambiar"), 400
    try:
        db.execute(f"UPDATE usuarios SET {', '.join(f'{c} = ?' for c in values)} "
                   f"WHERE id = ?", [*values.values(), uid])
        admins = db.execute("SELECT COUNT(*) FROM usuarios "
                            "WHERE rol = 'admin' AND activo = 1").fetchone()[0]
        if admins == 0:
            db.rollback()
            return jsonify(error="Tiene que quedar al menos un administrador activo"), 409
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        return jsonify(error="Ese nombre de usuario ya existe"), 409
    u = db.execute("SELECT * FROM usuarios WHERE id = ?", (uid,)).fetchone()
    if uid == current_user()["id"] and "password_hash" in values:
        session["ver"] = password_version(u["password_hash"])
    return jsonify(public_user(u))


# ---------------------------------------------------------------------------
# API genérica (CRUD)
# ---------------------------------------------------------------------------

@app.errorhandler(ValidationError)
def on_validation_error(e):
    return jsonify(error=str(e)), 400


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/api/meta")
def meta():
    db = get_db()
    campanas = {campana_de(datetime.date.today().isoformat())}
    for t in ("recolecciones", "entregas", "tratamientos"):
        for r in db.execute(f"SELECT DISTINCT campana(fecha) c FROM {t}"):
            campanas.add(r["c"])
    return jsonify(categorias=CATEGORIAS,
                   campanas=sorted(campanas, reverse=True),
                   campana_actual=campana_de(datetime.date.today().isoformat()))


@app.route("/api/<table>", methods=["GET"])
def list_rows(table):
    if table not in TABLES:
        return jsonify(error="No existe"), 404
    spec = TABLES[table]
    sql, args = f"SELECT * FROM {table}", []
    camp = request.args.get("campana")
    if camp and spec.get("campana"):
        sql += " WHERE campana(fecha) = ?"
        args.append(camp)
    sql += f" ORDER BY {spec['order']}"
    return jsonify(rows(get_db().execute(sql, args)))


@app.route("/api/<table>", methods=["POST"])
def create_row(table):
    if table not in TABLES:
        return jsonify(error="No existe"), 404
    data = request.get_json(force=True) or {}
    values = parse_payload(table, data)
    check_required(table, values)
    db = get_db()
    if table == "recolecciones" and not data.get("forzar"):
        activos = tratamientos_activos(db, values["sector_id"], values["fecha"])
        if activos:
            t = activos[0]
            return jsonify(
                error=(f"Plazo de seguridad en vigor: «{t['producto']}» aplicado "
                       f"el {t['fecha']}, no se puede recolectar hasta el "
                       f"{t['recolectable_desde']}."),
                plazo_seguridad=activos), 409
    cols = ", ".join(values)
    marks = ", ".join("?" for _ in values)
    try:
        cur = db.execute(f"INSERT INTO {table} ({cols}) VALUES ({marks})",
                         list(values.values()))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify(error="Referencia no válida"), 400
    return jsonify(dict(db.execute(f"SELECT * FROM {table} WHERE id = ?",
                                   (cur.lastrowid,)).fetchone())), 201


@app.route("/api/<table>/<int:row_id>", methods=["PUT"])
def update_row(table, row_id):
    if table not in TABLES:
        return jsonify(error="No existe"), 404
    db = get_db()
    current = db.execute(f"SELECT * FROM {table} WHERE id = ?", (row_id,)).fetchone()
    if current is None:
        return jsonify(error="No existe"), 404
    values = parse_payload(table, request.get_json(force=True) or {})
    check_required(table, {**dict(current), **values})
    if values:
        sets = ", ".join(f"{c} = ?" for c in values)
        try:
            db.execute(f"UPDATE {table} SET {sets} WHERE id = ?",
                       [*values.values(), row_id])
            db.commit()
        except sqlite3.IntegrityError:
            return jsonify(error="Referencia no válida"), 400
    return jsonify(dict(db.execute(f"SELECT * FROM {table} WHERE id = ?",
                                   (row_id,)).fetchone()))


@app.route("/api/<table>/<int:row_id>", methods=["DELETE"])
def delete_row(table, row_id):
    if table not in TABLES:
        return jsonify(error="No existe"), 404
    db = get_db()
    try:
        db.execute(f"DELETE FROM {table} WHERE id = ?", (row_id,))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify(error="No se puede borrar: tiene registros asociados"), 409
    return jsonify(ok=True)


# ---------------------------------------------------------------------------
# Resumen de campaña y plazos de seguridad
# ---------------------------------------------------------------------------

@app.route("/api/plazos")
def plazos():
    """Sectores con plazo de seguridad en vigor en la fecha dada (hoy)."""
    fecha = request.args.get("fecha") or datetime.date.today().isoformat()
    db = get_db()
    out = []
    for s in db.execute("SELECT id FROM sectores"):
        activos = tratamientos_activos(db, s["id"], fecha)
        if activos:
            out.append({"sector_id": s["id"], "tratamientos": activos,
                        "recolectable_desde": activos[0]["recolectable_desde"]})
    return jsonify(out)


@app.route("/api/resumen")
def resumen():
    db = get_db()
    camp = request.args.get("campana") or campana_de(datetime.date.today().isoformat())
    w = "campana(fecha) = ?"

    total = db.execute(
        f"SELECT COALESCE(SUM(kg), 0) kg, COALESCE(SUM(cajas), 0) cajas, "
        f"COUNT(DISTINCT fecha) dias FROM recolecciones WHERE {w}",
        (camp,)).fetchone()
    por_categoria = rows(db.execute(
        f"SELECT categoria, SUM(kg) kg FROM recolecciones WHERE {w} "
        f"GROUP BY categoria", (camp,)))
    # Semana empezando en lunes (strftime %w: 0 = domingo).
    por_semana = rows(db.execute(
        f"""SELECT date(fecha, '-' || ((CAST(strftime('%w', fecha) AS INTEGER) + 6) % 7)
                        || ' days') semana,
                   SUM(kg) kg
            FROM recolecciones WHERE {w} GROUP BY semana ORDER BY semana""",
        (camp,)))
    por_sector = rows(db.execute(
        f"""SELECT s.id sector_id, s.nombre sector, f.nombre finca, s.variedad,
                   s.superficie_ha, s.plantas, COALESCE(SUM(r.kg), 0) kg
            FROM sectores s JOIN fincas f ON f.id = s.finca_id
            LEFT JOIN recolecciones r ON r.sector_id = s.id AND campana(r.fecha) = ?
            GROUP BY s.id ORDER BY kg DESC""", (camp,)))
    for s in por_sector:
        s["kg_ha"] = s["kg"] / s["superficie_ha"] if s["superficie_ha"] else None
        s["g_planta"] = s["kg"] * 1000 / s["plantas"] if s["plantas"] else None
    por_trabajador = rows(db.execute(
        f"""SELECT t.id trabajador_id, t.nombre, t.cuadrilla,
                   SUM(r.kg) kg, SUM(r.cajas) cajas, COUNT(DISTINCT r.fecha) dias
            FROM recolecciones r JOIN trabajadores t ON t.id = r.trabajador_id
            WHERE campana(r.fecha) = ?
            GROUP BY t.id ORDER BY kg DESC""", (camp,)))
    for t in por_trabajador:
        t["kg_dia"] = t["kg"] / t["dias"] if t["dias"] else None
    ent = db.execute(
        f"SELECT COALESCE(SUM(kg), 0) kg, COALESCE(SUM(kg * precio_kg), 0) importe, "
        f"SUM(CASE WHEN precio_kg IS NOT NULL THEN kg END) kg_liquidados "
        f"FROM entregas WHERE {w}", (camp,)).fetchone()
    ent_cat = rows(db.execute(
        f"""SELECT categoria, SUM(kg) kg, SUM(kg * precio_kg) importe,
                   SUM(kg * precio_kg) / NULLIF(SUM(CASE WHEN precio_kg IS NOT NULL
                                                    THEN kg END), 0) precio_medio
            FROM entregas WHERE {w} GROUP BY categoria""", (camp,)))

    return jsonify(
        campana=camp,
        recolectado={"kg": total["kg"], "cajas": total["cajas"], "dias": total["dias"]},
        por_categoria=por_categoria,
        por_semana=por_semana,
        por_sector=por_sector,
        por_trabajador=por_trabajador,
        entregado={"kg": ent["kg"], "importe": ent["importe"],
                   "precio_medio": (ent["importe"] / ent["kg_liquidados"]
                                    if ent["kg_liquidados"] else None),
                   "por_categoria": ent_cat},
    )


# ---------------------------------------------------------------------------
# Exportación a Excel
# ---------------------------------------------------------------------------

def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/export/cuaderno")
def export_cuaderno():
    """Cuaderno de campo: registro de tratamientos fitosanitarios."""
    camp = request.args.get("campana") or campana_de(datetime.date.today().isoformat())
    data = get_db().execute(
        """SELECT t.fecha, f.nombre finca, s.nombre sector, s.variedad,
                  s.superficie_ha, t.plaga, t.producto, t.num_registro,
                  t.materia_activa, t.dosis, t.plazo_seguridad_dias,
                  date(t.fecha, '+' || t.plazo_seguridad_dias || ' days'), t.aplicador
           FROM tratamientos t JOIN sectores s ON s.id = t.sector_id
           JOIN fincas f ON f.id = s.finca_id
           WHERE campana(t.fecha) = ? ORDER BY t.fecha""", (camp,)).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tratamientos"
    ws.append(["Fecha", "Finca", "Sector", "Variedad", "Superficie (ha)",
               "Plaga / problema", "Producto", "Nº registro", "Materia activa",
               "Dosis", "Plazo seguridad (días)", "Recolectable desde",
               "Aplicador"])
    for r in data:
        ws.append(list(r))
    return xlsx_response(wb, f"cuaderno-campo-{camp.replace('/', '-')}.xlsx")


@app.route("/api/export/campana")
def export_campana():
    """Recolección y entregas de la campaña, una hoja por tipo."""
    camp = request.args.get("campana") or campana_de(datetime.date.today().isoformat())
    db = get_db()
    wb = Workbook()
    ws = wb.active
    ws.title = "Recolección"
    ws.append(["Fecha", "Finca", "Sector", "Variedad", "Trabajador",
               "Cuadrilla", "Categoría", "Cajas", "Kg"])
    for r in db.execute(
            """SELECT r.fecha, f.nombre, s.nombre, s.variedad, t.nombre,
                      t.cuadrilla, r.categoria, r.cajas, r.kg
               FROM recolecciones r JOIN sectores s ON s.id = r.sector_id
               JOIN fincas f ON f.id = s.finca_id
               LEFT JOIN trabajadores t ON t.id = r.trabajador_id
               WHERE campana(r.fecha) = ? ORDER BY r.fecha""", (camp,)):
        ws.append(list(r))
    ws = wb.create_sheet("Entregas")
    ws.append(["Fecha", "Destino", "Albarán", "Categoría", "Kg", "€/kg", "Importe"])
    for r in db.execute(
            """SELECT fecha, destino, albaran, categoria, kg, precio_kg,
                      kg * precio_kg FROM entregas
               WHERE campana(fecha) = ? ORDER BY fecha""", (camp,)):
        ws.append(list(r))
    return xlsx_response(wb, f"campana-{camp.replace('/', '-')}.xlsx")


def crear_admin(usuario):
    """Crea (o restablece) un administrador desde la consola."""
    password = getpass.getpass(f"Contraseña para {usuario}: ")
    if password != getpass.getpass("Repítela: "):
        sys.exit("Las contraseñas no coinciden.")
    try:
        check_password_rules(password)
    except ValidationError as e:
        sys.exit(str(e))
    with app.app_context():
        db = get_db()
        db.execute(
            """INSERT INTO usuarios (usuario, nombre, password_hash, rol)
               VALUES (?, ?, ?, 'admin')
               ON CONFLICT(usuario) DO UPDATE SET
                 password_hash = excluded.password_hash, rol = 'admin', activo = 1""",
            (usuario, usuario, generate_password_hash(password)))
        db.commit()
    print(f"Administrador «{usuario}» listo.")


init_db()

if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "crear-admin":
        crear_admin(sys.argv[2])
        sys.exit(0)
    if len(sys.argv) > 1:
        sys.exit("Uso: python server.py [crear-admin <usuario>]")
    with app.app_context():
        if not get_db().execute("SELECT 1 FROM usuarios").fetchone():
            print("Aviso: no hay usuarios. Crea uno con:\n"
                  "  python server.py crear-admin <usuario>")
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5001)), debug=False)
